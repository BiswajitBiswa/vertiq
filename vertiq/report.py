"""One command from exported tables to a written report: everything Vertiq can say about the data, in its
own words, and what it cannot say yet.

  python -m vertiq.report config.json [--out report.md]

config.json names the tables and how they map to the contract; nothing else is configured:

  {"records": {"file": "orders.csv", "mapping": {"entity": "Customer", "time": "Date", "clock": "Time",
               "value": "Total", "attrs": ["Order No", "Channel"], "date_format": "%d/%m/%Y"}},
   "items":   {"file": "order_lines.csv", "key": "Order No", "event_key": "Order No", "item": "Item",
               "qty": "Qty", "item_group": "Category"},
   "entity_table": {"file": "customers.csv", "trainable": "trainable"},
   "noun": "orders", "currency": "₹", "entities": "customers"}

"records" can also list several event files about the same entities, each named by its "event" (default: the
file name), with "target" naming the event to predict. "target" may also be an outcome as written in a column ("Fail" in a
results column); that column then gives each record's type:

  {"records": [{"file": "errors.csv", "event": "error"}, {"file": "failures.csv", "event": "failure"}],
   "entity_table": {"file": "machines.csv"}, "target": "failure", "noun": "events", "entities": "machines"}

Predictions try return windows of 30, 60 and 90 days ("windows": [...] or "window_days": N to change), with the
gate's confidence corrected for the number tried, and search the inputs vertiq.menu proposes from the data
("tune": false to skip). A column listing each record's items in one cell is read as line items when no items
table is given ("item_text": false to skip).

"scores": "scores.csv" writes every entity's prediction from each promoted model to that file, with the reasons
that raise or lower it, and lists the highest in the report.

Consent: with "trainable" naming a column of the entity table, only entities whose value there is true (true, 1,
yes) are used to train or test models; an entity missing from the table is not. "train": false keeps every entity
out of training. Counts, the summary and patterns still use every record: they describe, they do not learn.

CSV is read with the standard library; .xlsx needs openpyxl (pip install openpyxl). Every mapping is checked
against the rows by vertiq.mapping.repair before use, and every section below comes from a Vertiq module:
summary and narrate (what happened), patterns (tested patterns), affinity (items bought together) and the
tasks behind the evaluation gate (predictions, or why there are none yet).
"""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from . import affinity, harness, narrate, patterns, summary
from .mapping import Mapping, apply, attach_items, repair
from .tasks import predict_event, predict_value


def _grid(path: Path) -> list[list]:
    """All cells of the first sheet with data, as rows of values. The format is read from the file's first bytes,
    not its name, since exports are often misnamed (an .xls that is really .xlsx): a zip is Excel 2007+
    (needs openpyxl), an OLE container is Excel 97-2003 (needs xlrd), anything else is read as CSV."""
    head = path.open("rb").read(8)
    if head.startswith(b"PK"):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise SystemExit("reading Excel files needs openpyxl: pip install openpyxl (or save the sheet as CSV)") from e
        import io
        wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(any(v not in (None, "") for v in r) for r in rows):
                return rows
        return []
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        try:
            import xlrd
        except ImportError as e:
            raise SystemExit("reading Excel 97-2003 files needs xlrd: pip install xlrd (or save the sheet as CSV)") from e
        book = xlrd.open_workbook(str(path))
        for sh in book.sheets():
            if sh.nrows:
                return [sh.row_values(i) for i in range(sh.nrows)]
        return []
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(f, dialect)]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d %H:%M:%S") if (getattr(v, "hour", 0) or getattr(v, "minute", 0)) else v.strftime("%Y-%m-%d")
    return str(v)


def read_table(path: str | Path) -> list[dict]:
    """A table as {column: text} rows. The header is the first row, among the first 20, that fills at least 80% of
    the widest row with text: title lines and blank rows above it are skipped. Repeated or empty header names
    are made unique ("column 7", "Total (2)")."""
    rows = [[_cell(v) for v in r] for r in _grid(Path(path))]
    if not rows:
        return []
    width = max(sum(1 for v in r if v.strip()) for r in rows[:20])
    at = next(i for i, r in enumerate(rows[:20]) if sum(1 for v in r if v.strip()) >= 0.8 * width
              and all(not v.strip() or not v.replace(".", "").replace("-", "").isdigit() for v in r))
    head, seen = [], Counter()
    for i, h in enumerate(rows[at]):
        name = h.strip() or f"column {i + 1}"
        seen[name] += 1
        head.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return [dict(zip(head, r + [""] * (len(head) - len(r)))) for r in rows[at + 1:]]


def _build_many(config: dict, base: Path) -> tuple[object, dict, list[str]]:
    """Several event files about the same entities (a machine's error log and its failure records, a customer's
    orders and returns): each file is read and mapped on its own, its events named after the file (or its
    "event"), and the files merged on the entity; an entity table joins once, at the end."""
    from .contract import Dataset, Entity
    ents: dict[str, Entity] = {}
    events, groups, log = [], {}, []
    for r in config["records"]:
        name = r.get("event") or Path(r["file"]).stem
        sub = {**config, "records": r, "entity_table": None, "_entity_table": config.get("entity_table"), "event": name}
        ds, g, lg = build(sub, base)
        log += [f"{name}: {x}" for x in lg]
        if ds is None:
            continue
        for e in ds.entities:
            old = ents.get(e.id)
            ents[e.id] = e if old is None else Entity(e.id, old.group or e.group, {**(e.attrs or {}), **(old.attrs or {})})
        events += ds.events
        groups.update(g)
    if not events:
        return None, {}, log
    ds = Dataset([ents[k] for k in sorted(ents)], sorted(events, key=lambda e: e.at))
    log.append(f"Merged: {len(events):,} events of {len({e.type for e in events})} types from {len(config['records'])} files")
    if config.get("entity_table"):
        from .tables import attach_entities
        et = config["entity_table"]
        ds, more, _labels = attach_entities(ds, read_table(base / et["file"]), et.get("key"))
        log += [f"Entities: {n}" for n in more]
    return ds, groups, log


def build(config: dict, base: Path = Path(".")) -> tuple[object, dict, list[str]]:
    """The dataset, item groups and loading notes from a config. Without a mapping, Vertiq proposes one from
    the values (vertiq.tables.propose); without join columns, it finds them (vertiq.tables.find_join).
    "records" may also be a list of event files about the same entities (see _build_many)."""
    from dataclasses import replace
    if isinstance(config["records"], list):
        return _build_many(config, base)
    from .tables import find_join, find_lines, propose, records_from_lines
    rec = config["records"]
    rows = read_table(base / rec["file"])
    log = []
    it = dict(config.get("items") or {})
    lines = read_table(base / it["file"]) if it else []
    if not it and not rec.get("mapping"):
        spec, why = find_lines(rows)                 # a single table of line items: build the records from it
        if spec:
            log += [f"Table shape: {w}" for w in why]
            lines = rows
            rows = records_from_lines(rows, spec)
            it = {"key": spec["key"], "event_key": spec["key"], "item": spec["item"], "qty": spec.get("qty"),
                  "item_group": spec.get("item_group")}
    if it and not it.get("key"):
        found, why = find_join(rows, lines)          # first: the join key names the record, not whom it is about
        log += [f"Line items: {w}" for w in why]
        if found:
            it.update(found)
    from .tables import entity_table
    shape = entity_table(rows) if not it else None
    if shape and shape["counts"]:
        log.append(f"Table shape: one row per entity (key {shape['key']!r}) with totals {shape['counts']}; each row is "
                   "read as the entity's latest activity, and its totals describe, never predict")
    if rec.get("mapping"):
        proposal = Mapping.parse(rec["mapping"])
    else:
        hint = None
        et = config.get("entity_table") or config.get("_entity_table")
        if et:                                        # the entity table names the entity: its key's column
            from .tables import entity_key_hint
            hint = entity_key_hint(rows, read_table(base / et["file"]))
        proposal, why = propose(rows, record_key=it.get("event_key"), latest=bool(shape and shape["counts"]),
                                entity_hint=hint)
        log += [f"Mapping proposal: {w}" for w in why]
        if proposal is None:
            return None, {}, log
    m, notes = repair(proposal, rows, exclude={it["event_key"]} if it.get("event_key") else ())
    ds, problems = apply(m, rows)
    target = config.get("target")
    if target and ds.events and target not in {e.type for e in ds.events}:
        # the outcome is named as written in the data ("Fail" in a results column): that column is the event type
        holder = next((a for a in m.attrs if any(str((e.attrs or {}).get(a, "")).strip() == target for e in ds.events)), None)
        if holder:
            ds = type(ds)(ds.entities, [replace(e, type=str((e.attrs or {}).get(holder) or "").strip() or "other",
                                                attrs={k: v for k, v in (e.attrs or {}).items() if k != holder})
                                        for e in ds.events], ds.actions)
            log.append(f"Target: {target!r} is a value of {holder!r}, so {holder!r} gives each record's type")
    if ds.events and not any(e.type for e in ds.events):
        noun = config.get("noun", "")
        name = config.get("event") or (noun[:-1] if noun.endswith("s") else noun) or "record"
        ds = type(ds)(ds.entities, [replace(e, type=name) for e in ds.events], ds.actions)
    log += [f"Mapping used: {json.dumps(m.as_dict())}"] + [f"Mapping check: {n}" for n in notes]
    log += [f"Loading: {p}" for p in problems[:12]] + ([f"Loading: and {len(problems) - 12} more"] if len(problems) > 12 else [])
    groups: dict = {}
    if config.get("entity_table"):                    # a customer list, asset register or policy master
        from .tables import attach_entities
        et = config["entity_table"]
        ds, more, _labels = attach_entities(ds, read_table(base / et["file"]), et.get("key"))
        log += [f"Entities: {n}" for n in more]
    if it:
        if it.get("key"):
            ds, groups, more = attach_items(ds, lines, key=it["key"], event_key=it["event_key"], item=it["item"],
                                            qty=it.get("qty"), item_group=it.get("item_group"))
            log += [f"Line items: {n}" for n in more]
    elif config.get("item_text", True):
        # no line-item table: a column listing each record's items in one cell ("Shirtx999,Tshirtx699") stands in
        from .tables import _records, find_item_text, lines_from_text
        col, why = find_item_text(rows, exclude={m.entity, m.time, m.value or ""})
        data = _records(rows)
        key = next((a for a in m.attrs if a != col and len({str(r.get(a) or "").strip() for r in data}) == len(data)), None)
        if col and key:
            ds, groups, more = attach_items(ds, lines_from_text(rows, col, key), key="key", event_key=key, item="item")
            log += [f"Line items: {w}" for w in why] + [f"Line items: {n}" for n in more]
    return ds, groups, log


def run(config: dict, base: Path = Path("."), collect: dict | None = None) -> str:
    """The report as Markdown. `collect`, when given, is filled with the same results as data (loading notes,
    counts, patterns, combinations, and each prediction task's gate report) for benchmarks to record."""
    collect = collect if collect is not None else {}
    ds, groups, log = build(config, base)
    collect["log"] = log
    noun, cur = config.get("noun", "records"), config.get("currency", "")
    if ds is None:
        return "\n".join(["# Vertiq report", "", "## Loading", ""] + [f"- {x}" for x in log] + ["", "No mapping could be proposed."]) + "\n"
    collect.update(records=len(ds.events), entities=len(ds.entities))
    out = ["# Vertiq report", "", f"{len(ds.events)} {noun} from {len(ds.entities)} entities.", "", "## Loading", ""]
    out += [f"- {x}" for x in log]
    if not ds.events:
        return "\n".join(out + ["", "Nothing loaded; see the notes above."]) + "\n"
    span = max(1, (max(e.at for e in ds.events).date() - min(e.at for e in ds.events).date()).days + 1)
    days = config.get("summary_days") or max(1, span // 2)
    out += ["", "## What happened", "", narrate.summary(summary.summarise(ds, days=days), entities=config.get("entities", "entities"))]
    pat = patterns.find(ds, item_groups=groups or None, noun=noun, currency=cur)
    collect["patterns"] = {k: pat.get(k) for k in ("status", "features", "tested")} | {"found": [f["text"] for f in pat.get("found", [])]}
    out += ["", "## Patterns", ""]
    if pat["status"] == "learning":
        out.append(f"Not enough {noun} to test patterns: {pat['records']}, {pat['need']} needed.")
    else:
        out.append(f"{pat['features']} features found in the data; {pat['tested']} comparisons tested; "
                   f"{len(pat['found'])} hold after false-discovery control at {pat['q']:.0%} and are large enough to matter.")
        out += [f"- {f['text']}" for f in pat["found"]]
    lines = [(e.entity, f"{e.entity}|{e.at.isoformat()}|{i}", e.at, name) for i, e in enumerate(ds.events) for name, _q in (e.items or ())]
    if lines:
        a = affinity.analyse(lines)
        combos = a.get("combos", [])
        collect["combos"] = [{"items": c["items"], "bills_together": c["bills_together"], "lift": c["lift"]} for c in combos[:10]]
        out += ["", "## Items bought together", ""]
        out += ([f"- {' + '.join(c['items'])}: on {c['bills_together']} {noun} together, {c['lift']} times as often as chance" for c in combos[:10]]
                or [f"No pair of items appears together on at least {affinity.MIN_PAIR} {noun}, the minimum for a combination."])
    types = sorted({e.type for e in ds.events if e.type})
    target = config.get("target") or (types[0] if len(types) == 1 else None)
    out += ["", "## Predictions", ""]
    allowed, why = trainable(config, ds)
    if allowed is not None:
        from .contract import Dataset
        if not allowed:
            out.append(f"- No model was trained: none of the {len(ds.entities):,} entities may be used for training ({why}). "
                       "The sections above describe the data; they do not train on it.")
            return "\n".join(out) + "\n"
        out.append(f"- Models use only the {len(allowed):,} of {len(ds.entities):,} entities that may be used for training ({why}).")
        ds = Dataset([e for e in ds.entities if e.id in allowed], [e for e in ds.events if e.entity in allowed], ds.actions)
    tasks = []
    if target:
        windows = config.get("windows") or ([config["window_days"]] if config.get("window_days") else [30, 60, 90])
        # several windows are several chances to pass the gate: correct the confidence for how many are tried
        conf = 1 - (1 - harness.CONFIDENCE) / len(windows)
        if len(windows) > 1:
            out.append(f"- Windows tried: {', '.join(f'{w} days' for w in windows)}; each must pass the gate at "
                       f"{conf:.1%} confidence, corrected for trying {len(windows)}.")
        tune = config.get("tune", TUNE)
        for w in windows:
            tasks.append(_task(ds, "event", target, w, conf, tune, out))
        if any(e.value for e in ds.events):
            tasks.append(_task(ds, "value", target, max(windows), harness.CONFIDENCE, tune, out))
    else:
        out.append("- No single event type to predict; set \"target\" in the config.")
    collect["predictions"] = [t.report() for t in tasks]
    if config.get("scores"):
        out += scores(tasks, base / config["scores"], cur, config.get("entities", "entities"))
    return "\n".join(out) + "\n"


TUNE = True         # search candidate inputs by default (tests that only check wording turn it off)


def _task(ds, kind: str, target: str, window: int, confidence: float, tune: bool, out: list[str]):
    """One prediction task through the gate, written to `out`. With `tune`, vertiq.autotune first tries the
    candidate inputs vertiq.menu proposes from the data (counts, shares and recency by attribute value and by
    item), keeping each only if it improves validation folds of the training entities; the tuned model must
    then also beat the untuned one on the held-out entities."""
    fn = predict_event if kind == "event" else predict_value
    if not tune:
        t = fn(ds, target=target, window_days=window, confidence=confidence)
        out.append("- " + narrate.gate(t.report()))
        return t
    from . import autotune, menu
    specs = menu.propose(ds, target=target, window_days=window, limit=12)
    res = autotune.search(ds, target=target, window_days=window, kind=kind, features=menu.features(specs, ds),
                          rounds=1, confidence=confidence)
    t = res["task"]
    line = "- " + narrate.gate(t.report())
    if res["features"]:
        names = ", ".join(menu.words(f) for f in res["features"])
        line += (f" Inputs added by the search ({len(specs)} tried): {names}; "
                 + ("the tuned model was kept." if res["tuned_promoted"] else "the tuned model did not beat the untuned one."))
    else:
        line += f" None of the {len(specs)} candidate inputs the search tried improved it."
    out.append(line)
    return t


def scores(tasks: list, path: Path, cur: str, entities: str, top: int = 10) -> list[str]:
    """Every entity's prediction from each promoted model, written to `path` as CSV (entity, model, prediction,
    what raises it, what lowers it), and the highest ones as report lines. Models the gate kept give none: their
    predictions did not beat the rule in use. One-off entities (no id) are left out."""
    promoted = [t for t in tasks if t.status == "trained"]
    lines = ["", "## Scores", ""]
    if not promoted:
        return lines + ["No model was promoted, so there are no scores: the rules in use stay."]
    rows = []
    for t in promoted:
        for eid in t.b.by:
            if eid.startswith("anonymous:"):
                continue
            sc = t.score(eid)
            if sc:
                what = (f"chance of {t.b.target} within {t.window} days" if t.kind == "event" else f"value of the next {t.b.target}")
                rows.append({"entity": eid, "model": what, "prediction": sc.get("chance", sc.get("value")),
                             "raises": "; ".join(sc["raises"]), "lowers": "; ".join(sc["lowers"])})
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["entity", "model", "prediction", "raises", "lowers"])
        w.writeheader()
        w.writerows(rows)
    lines.append(f"{len(rows):,} predictions for {len({r['entity'] for r in rows}):,} {entities}, written to {path.name}. "
                 f"Highest {min(top, len(rows))} per model:")
    for what in dict.fromkeys(r["model"] for r in rows):
        lines += ["", f"**{what.capitalize()}**", ""]
        best = sorted((r for r in rows if r["model"] == what), key=lambda r: -r["prediction"])[:top]
        for r in best:
            val = f"{r['prediction']:.0%}" if what.startswith("chance") else f"{cur}{r['prediction']:,.0f}"
            why = "; ".join(x for x in (f"raised by {r['raises']}" if r["raises"] else "",
                                        f"lowered by {r['lowers']}" if r["lowers"] else "") if x)
            lines.append(f"- {r['entity'][:12]}: {val}" + (f" ({why})" if why else ""))
    return lines


def trainable(config: dict, ds) -> tuple[set | None, str]:
    """The entities that may be used to train models, and why; (None, "") when the config says nothing."""
    if config.get("train") is False:
        return set(), '"train": false in the config'
    col = (config.get("entity_table") or {}).get("trainable")
    if not col:
        return None, ""
    yes = {e.id for e in ds.entities if str((e.attrs or {}).get(col, "")).strip().lower() in ("true", "1", "yes", "y")}
    return yes, f"column {col!r} of the entity table"


def main():
    ap = argparse.ArgumentParser(description="A written report from exported tables (see the module docstring).")
    ap.add_argument("config")
    ap.add_argument("--out")
    a = ap.parse_args()
    cfg = Path(a.config)
    text = run(json.loads(cfg.read_text()), cfg.parent)
    if a.out:
        Path(a.out).write_text(text)
    print(text)


if __name__ == "__main__":
    main()

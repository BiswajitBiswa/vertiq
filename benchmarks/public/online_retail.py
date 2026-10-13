"""Online Retail II benchmark (proposal 0001, part C, protocol 15.1): repeat purchase on real transactions.

Data: UCI Machine Learning Repository, "Online Retail II" (CC BY 4.0). Download it yourself and pass the
path; nothing is downloaded by this script and no data is committed. Accepts the .xlsx file (both
sheets, needs openpyxl) or a CSV export with the same columns:
Invoice, StockCode, Description, Quantity, InvoiceDate, Price, Customer ID, Country.

Protocol: drop cancellations (invoices starting with "C") and lines without a customer; aggregate lines
to one order per invoice (value = sum of quantity x price); predict whether each customer orders again
within 60 days, with the standard cutoffs and entity-level holdout; compare the recency baseline with
the linear, boosted and auto learners, with and without the survival feature. scikit-learn and LightGBM
are used as external references only when installed, on the same features and split.

Run:  python benchmarks/public/online_retail.py --data path/to/online_retail_II.xlsx
Writes benchmarks/results/public.json.
"""
from __future__ import annotations

import argparse
import csv
import json
import platform
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from vertiq.contract import Dataset, Entity, Event  # noqa: E402
from vertiq.tasks import predict_event  # noqa: E402

DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%m/%d/%Y %H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M")


def _date(v) -> datetime | None:
    if isinstance(v, datetime):
        return v
    for f in DATE_FORMATS:
        try:
            return datetime.strptime(str(v).strip(), f)
        except ValueError:
            continue
    return None


def read_lines(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook          # benchmark-only dependency
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = []
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            head = [str(h).strip() for h in next(it)]
            rows += [dict(zip(head, r)) for r in it]
        return rows
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        return list(csv.DictReader(f))


def to_dataset(lines: list[dict]) -> tuple[Dataset, dict]:
    """Invoice lines to one order event per invoice, per customer. Returns the dataset and counts."""
    orders: dict[str, dict] = {}
    dropped = {"cancelled": 0, "no_customer": 0, "bad_date": 0}
    for r in lines:
        inv = str(r.get("Invoice") or r.get("InvoiceNo") or "").strip()
        cust = str(r.get("Customer ID") or r.get("CustomerID") or "").strip()
        if inv.upper().startswith("C"):
            dropped["cancelled"] += 1
            continue
        if not cust or cust.lower() == "none":
            dropped["no_customer"] += 1
            continue
        at = _date(r.get("InvoiceDate"))
        if at is None:
            dropped["bad_date"] += 1
            continue
        try:
            amount = float(r.get("Quantity") or 0) * float(r.get("Price") or r.get("UnitPrice") or 0)
        except ValueError:
            amount = 0.0
        o = orders.setdefault(inv, {"customer": cust.split(".")[0], "at": at, "value": 0.0,
                                    "country": str(r.get("Country") or ""), "items": []})
        o["value"] += amount
        o["items"].append(str(r.get("StockCode") or ""))
    countries = {}
    events = []
    for o in orders.values():
        countries.setdefault(o["customer"], o["country"])
        events.append(Event(o["customer"], o["at"], "order", round(o["value"], 2), {"Country": o["country"]},
                            tuple((i, 1) for i in o["items"])))
    ents = [Entity(c, group=g) for c, g in countries.items()]
    return Dataset(ents, events), {"lines": len(lines), "orders": len(orders), "customers": len(ents), "dropped": dropped}


def _external(task) -> dict:
    """scikit-learn and LightGBM on the same features and entity split, when installed."""
    out = {}
    try:
        from sklearn.linear_model import LogisticRegression
        from sklearn.metrics import roc_auc_score
    except ImportError:
        return {"note": "scikit-learn not installed"}
    from vertiq import harness
    b = task.b
    rows, y, keys = [], [], []
    for cut in b.cutoffs(task.window, 4):
        cad = b.cadence(cut, task.window)
        for eid, evs in b.by.items():
            f = b.row(eid, cut, cad)
            if f is None:
                continue
            rows.append(f)
            y.append(1 if any(e.type == b.target and cut <= e.at < cut + timedelta(days=task.window) for e in evs) else 0)
            keys.append(eid)
    names = sorted({k for r in rows for k in r})
    X = [[r.get(k, 0.0) for k in names] for r in rows]
    test = [harness.held_out(k) for k in keys]
    tr = [i for i in range(len(y)) if not test[i]]
    te = [i for i in range(len(y)) if test[i]]
    lr = LogisticRegression(max_iter=2000).fit([X[i] for i in tr], [y[i] for i in tr])
    out["sklearn_logistic_auc"] = round(roc_auc_score([y[i] for i in te], lr.predict_proba([X[i] for i in te])[:, 1]), 4)
    try:
        import lightgbm as lgb
        g = lgb.LGBMClassifier(n_estimators=300, learning_rate=0.05, verbose=-1).fit([X[i] for i in tr], [y[i] for i in tr])
        out["lightgbm_auc"] = round(roc_auc_score([y[i] for i in te], g.predict_proba([X[i] for i in te])[:, 1]), 4)
    except ImportError:
        out["lightgbm_auc"] = "not installed"
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--window", type=int, default=60)
    ap.add_argument("--out", default=str(ROOT / "benchmarks" / "results" / "public.json"))
    a = ap.parse_args()
    t0 = time.time()
    ds, counts = to_dataset(read_lines(Path(a.data)))
    prep = round(time.time() - t0, 1)
    runs = {}
    for name, kw in (("linear", {"learner": "logistic"}), ("boosted", {"learner": "boosted"}), ("auto", {"learner": "auto"}),
                     ("linear_survival", {"learner": "logistic", "survival": True})):
        t1 = time.time()
        task = predict_event(ds, target="order", window_days=a.window, **kw)
        m = task.report().get("metrics", {})
        runs[name] = {"status": task.status, "auc": m.get("auc"), "baseline_auc": m.get("baseline_auc"), "brier": m.get("brier"),
                      "intervals": m.get("intervals"), "learner": m.get("learner"), "seconds": round(time.time() - t1, 1)}
        print(name, runs[name])
        if name == "linear":
            runs["external"] = _external(task)
            print("external", runs["external"])
    doc = {"dataset": "Online Retail II (UCI)", "generated": datetime.utcnow().strftime("%Y-%m-%d"), "window_days": a.window,
           "counts": counts, "prep_seconds": prep, "python": platform.python_version(), "machine": platform.machine(), "runs": runs}
    Path(a.out).write_text(json.dumps(doc, indent=2, default=str))
    print("wrote", a.out)


if __name__ == "__main__":
    main()
